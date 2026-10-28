"""Read a resume (PDF / DOCX / TXT) and pull out skills and years of experience."""
from __future__ import annotations

import re
from pathlib import Path

# Canonical skill -> extra spellings. Extend freely (PRs welcome).
SKILLS: dict[str, list[str]] = {
    "Python": [], "Java": [], "JavaScript": ["js"], "TypeScript": ["ts"], "Go": ["golang"],
    "C++": ["cpp"], "C#": ["csharp", ".net", "dotnet"], "Rust": [], "Kotlin": [], "Swift": [],
    "PHP": [], "Ruby": ["rails", "ruby on rails"], "Scala": [], "SQL": [], "Bash": ["shell scripting"],
    "React": ["react.js", "reactjs"], "Next.js": ["nextjs"], "Angular": [], "Vue": ["vue.js", "vuejs"],
    "Node.js": ["node", "nodejs", "express", "express.js"], "Django": [], "Flask": [], "FastAPI": [],
    "Spring Boot": ["spring"], "HTML": [], "CSS": ["tailwind", "sass"], "Redux": [],
    "React Native": [], "Flutter": [], "Android": [], "iOS": [],
    "PostgreSQL": ["postgres"], "MySQL": [], "MongoDB": ["mongo"], "Redis": [], "Elasticsearch": [],
    "Kafka": [], "RabbitMQ": [], "GraphQL": [], "REST": ["rest api", "restful"], "gRPC": [],
    "AWS": ["amazon web services", "ec2", "s3", "lambda"], "GCP": ["google cloud"], "Azure": [],
    "Docker": [], "Kubernetes": ["k8s"], "Terraform": [], "CI/CD": ["github actions", "jenkins", "gitlab ci"],
    "Linux": [], "Git": [], "Microservices": [],
    "Machine Learning": ["ml"], "Deep Learning": [], "PyTorch": [], "TensorFlow": [], "scikit-learn": ["sklearn"],
    "NLP": [], "Computer Vision": [], "LLMs": ["llm", "large language models", "genai", "generative ai"],
    "RAG": ["retrieval augmented generation"], "LangChain": [], "LlamaIndex": [], "CrewAI": [],
    "OpenAI API": ["openai"], "Pandas": [], "NumPy": [], "Spark": ["pyspark"], "Airflow": [],
    "Data Engineering": ["etl"], "Power BI": [], "Tableau": [], "Excel": [],
    "Selenium": [], "Playwright": [], "Testing": ["pytest", "jest", "unit testing"],
    "Figma": [], "Product Management": [], "Salesforce": [], "SAP": [],
}


def _pattern(term: str) -> re.Pattern:
    esc = re.escape(term.lower())
    return re.compile(rf"(?<![a-z0-9+#.]){esc}(?![a-z0-9+#])", re.I)


_COMPILED = {canon: [_pattern(canon)] + [_pattern(a) for a in aliases] for canon, aliases in SKILLS.items()}


def skill_patterns(skill: str) -> list[re.Pattern]:
    return _COMPILED.get(skill) or [_pattern(skill)]


def find_skills(text: str, vocabulary: list[str] | None = None) -> list[str]:
    vocab = vocabulary or list(SKILLS)
    found = []
    for s in vocab:
        if any(p.search(text) for p in skill_patterns(s)):
            found.append(s)
    # "Go" is noisy in prose; keep it only if written as Golang or in a skills list context
    if "Go" in found and not re.search(r"(?i)\bgolang\b|\bgo\s*(,|/|\|)|(,|/|\|)\s*go\b", text):
        found.remove("Go")
    return found


def read_text(path: str | Path) -> str:
    p = Path(path)
    suffix = p.suffix.lower()
    if suffix == ".pdf":
        from pypdf import PdfReader
        return "\n".join((page.extract_text() or "") for page in PdfReader(str(p)).pages)
    if suffix == ".docx":
        import docx
        d = docx.Document(str(p))
        parts = [para.text for para in d.paragraphs]
        for t in d.tables:
            for row in t.rows:
                parts.extend(c.text for c in row.cells)
        return "\n".join(parts)
    return p.read_text(encoding="utf-8", errors="ignore")


def guess_years(text: str) -> float:
    m = re.findall(r"(\d{1,2}(?:\.\d)?)\s*\+?\s*(?:years?|yrs?)\s+(?:of\s+)?(?:professional\s+|industry\s+|work\s+)?experience", text, re.I)
    return max((float(x) for x in m), default=0.0)


def extract_profile(text: str) -> dict:
    return {"skills": find_skills(text), "years": guess_years(text)}
